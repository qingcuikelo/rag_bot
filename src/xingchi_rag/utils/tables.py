"""表格读取工具（CSV / TSV / XLSX），统一为记录列表。"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def read_csv_rows(path: str | Path) -> list[dict[str, str]]:
    """读取 CSV（UTF-8 BOM 兼容）。"""
    with Path(path).open("r", encoding="utf-8-sig", newline="") as fh:
        return [dict(r) for r in csv.DictReader(fh)]


def read_tsv_rows(path: str | Path) -> list[dict[str, str]]:
    """读取 TSV。"""
    with Path(path).open("r", encoding="utf-8-sig", newline="") as fh:
        return [dict(r) for r in csv.DictReader(fh, delimiter="\t")]


def read_sheet(path: str | Path, sheet_name: str) -> list[dict[str, Any]]:
    """读取指定 sheet，表头为键，跳过空行。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [("" if c is None else str(c).strip()) for c in rows[0]]
        records: list[dict[str, Any]] = []
        for row in rows[1:]:
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            records.append({h: v for h, v in zip(header, row, strict=False) if h})
        return records
    finally:
        wb.close()


def sheet_names(path: str | Path) -> list[str]:
    """返回工作簿全部 sheet 名。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
